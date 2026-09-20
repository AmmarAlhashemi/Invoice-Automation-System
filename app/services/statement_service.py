import os
import pandas as pd
import openpyxl
from openpyxl import load_workbook
from copy import copy
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx2pdf import convert
from config import Config
from app.services.invoice_service import refresh_excel_formulas

def set_cell_background(cell, fill_color):
    """Change the background color of a Word table cell"""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def format_cell_text(cell, text, font_name="Calibri", size=14, bold=False, align="center"):
    """Format text inside a Word table cell"""
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align == "center":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "left":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
    run = paragraph.add_run(str(text))
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    paragraph.paragraph_format.space_after = Pt(0)

def process_advance_payment(date_str, amount, description="ADVANCE RECEIVED"):
    """Process and add an advance payment in Excel"""
    DATE_COL = 1
    DESC_COL = 3
    AMOUNT_COL = 5
    
    try:
        wb = load_workbook(Config.EXCEL_PATH)
        ws = wb.active
        entry_date = datetime.strptime(date_str, Config.DATE_FORMAT)
        
        opening_rows = [r[0].row for r in ws.iter_rows(min_col=DESC_COL, max_col=DESC_COL) if str(r[0].value).strip() == Config.OPENING_TEXT]
        if not opening_rows:
            return False, "No opening balance row found."
            
        template_row = opening_rows[-2] - 3 if len(opening_rows) >= 2 else opening_rows[-1] - 3
        first_date_in_last_block = None
        
        for r in range(opening_rows[-1] + 1, ws.max_row + 1):
            val = ws.cell(r, DATE_COL).value
            if val:
                try:
                    first_date_in_last_block = datetime.strptime(str(val).strip(), Config.DATE_FORMAT)
                    break
                except:
                    continue
                    
        if first_date_in_last_block and entry_date < first_date_in_last_block and len(opening_rows) >= 2:
            search_start = opening_rows[-2]
            search_end = opening_rows[-1] - 1
        else:
            search_start = opening_rows[-1]
            search_end = ws.max_row
            
        insert_row = None
        for r in range(search_start + 1, search_end + 1):
            cell_val = ws.cell(r, DATE_COL).value
            if not cell_val or str(cell_val).strip() == "":
                insert_row = r
                break
            try:
                row_date = datetime.strptime(str(cell_val).strip(), Config.DATE_FORMAT)
                if row_date > entry_date:
                    insert_row = r
                    break
            except:
                if insert_row is None:
                    insert_row = r
                    break
                    
        if not insert_row:
            insert_row = search_end
            
        ws.insert_rows(insert_row)
        ws.row_dimensions[ws.max_row-1].height = 18.75
        
        if insert_row < opening_rows[-1]:
            opening_rows[-1] += 1
            
        for col in range(1, ws.max_column + 1):
            src = ws.cell(template_row, col)
            tgt = ws.cell(insert_row, col)
            if src.has_style:
                tgt.font = copy(src.font)
                tgt.border = copy(src.border)
                tgt.fill = copy(src.fill)
                tgt.alignment = copy(src.alignment)
                tgt.number_format = src.number_format
                tgt.protection = copy(src.protection)
                
        ws.cell(insert_row, DATE_COL).value = date_str
        ws.cell(insert_row, DESC_COL).value = description
        ws.cell(insert_row, AMOUNT_COL).value = amount
        
        final_total_row = ws.max_row
        previous_total_row = opening_rows[-1] - 1
        
        ws.cell(row=final_total_row, column=4).value = f"=SUM(D{opening_rows[-1]}:D{final_total_row-1})"
        ws.cell(row=final_total_row, column=5).value = f"=SUM(E{opening_rows[-1]}:E{final_total_row-1})"
        ws.cell(row=final_total_row, column=6).value = f"=F{opening_rows[-1]}+D{final_total_row}-E{final_total_row}"
        ws.cell(row=final_total_row, column=9).value = f"=SUM(I8:I{final_total_row-1})"
        ws.cell(row=final_total_row, column=11).value = f"=F{final_total_row}+I{final_total_row}"
        
        if insert_row < opening_rows[-1]:
            ws.cell(row=opening_rows[-1], column=6).value = f"=F{opening_rows[-1]-1}"
            ws.cell(row=previous_total_row, column=4).value = f"=SUM(D{opening_rows[-2]}:D{previous_total_row-1})"
            ws.cell(row=previous_total_row, column=5).value = f"=SUM(E{opening_rows[-2]}:E{previous_total_row-1})"
            ws.cell(row=previous_total_row, column=6).value = f"=F{opening_rows[-2]}+D{previous_total_row}-E{previous_total_row}"
            
        wb.save(Config.EXCEL_PATH)
        wb.close()
        refresh_excel_formulas()
        return True, "Advance payment added successfully."
    except Exception as e:
        return False, f"Excel Error: {e}"

def generate_statement_docs():
    """Generate account statement and export as Word and PDF files"""
    try:
        df = pd.read_excel(Config.EXCEL_PATH, header=None, engine='openpyxl')
        for col_idx in [3, 4, 7, 8]:
            df[col_idx] = pd.to_numeric(df[col_idx], errors='coerce').fillna(0)
            
        opening_indices = df[df[2] == Config.OPENING_TEXT].index.tolist()
        if not opening_indices:
            return False, "No opening balance found in statement."
            
        last_idx = opening_indices[-1]
        block_df = df.iloc[last_idx:].copy()
        
        bf_date_raw = df.iloc[last_idx - 3, 0]
        bf_date_str = pd.to_datetime(bf_date_raw).strftime('%m/%d/%Y') if pd.notnull(bf_date_raw) else str(bf_date_raw)
        calculated_bf = float(df.iloc[last_idx, 5]) + df[8].sum() - block_df[8].sum()
        excel_expected_balance = pd.to_numeric(block_df.iloc[-1, 10], errors='coerce')
        
        invoices, transfers = [], []
        last_inv_date_val = None
        
        for _, row in block_df.iterrows():
            desc = str(row[2])
            if "Attached Invoice" in desc:
                amt = float(row[7]) if str(row[6]).strip().lower() == "edited" else float(row[3])
                date_txt = pd.to_datetime(row[0]).strftime('%m/%d/%Y') if pd.notnull(row[0]) else str(row[0])
                last_inv_date_val = row[0]
                invoices.append({'date': date_txt, 'no': str(row[1]), 'amount': amt})
            elif "ADVANCE RECEIVED" in desc:
                date_val = row[0]
                date_txt = pd.to_datetime(date_val).strftime('%m/%d/%Y') if pd.notnull(date_val) else ""
                transfers.append({'date': date_txt, 'amount': float(row[4])})
                
        final_balance = sum(i['amount'] for i in invoices) + calculated_bf - sum(t['amount'] for t in transfers)
        
        file_date_part = "Unknown"
        if last_inv_date_val:
            try:
                file_date_part = pd.to_datetime(last_inv_date_val).strftime("%d-%m-%Y")
            except:
                file_date_part = str(last_inv_date_val).replace('/', '-').replace(' ', '-')
                
        output_name = f"Statement of AC - {file_date_part}"
        doc = Document(str(Config.WORD_STATEMENT_TEMPLATE))
        
        top_date = invoices[-1]['date'] if invoices else ""
        for p in doc.paragraphs:
            if "Date:" in p.text:
                p.text = ""
                run = p.add_run(f"Date: {top_date}")
                run.font.name = "Bell MT"
                run.font.size = Pt(24)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0)
                
        table = doc.tables[0]
        for inv in invoices:
            row = table.add_row()
            format_cell_text(row.cells[0], inv['date'], align="center")
            format_cell_text(row.cells[1], inv['no'], align="center")
            format_cell_text(row.cells[2], f"{inv['amount']:,.2f}", align="right")
            
        def add_summary_row(label, value):
            row = table.add_row()
            for i in range(3):
                set_cell_background(row.cells[i], "DEEAF6")
            row.cells[0].merge(row.cells[1])
            format_cell_text(row.cells[0], label, bold=True, align="left")
            format_cell_text(row.cells[2], f"{value:,.2f}", size=16, bold=True, align="right")
            
        add_summary_row("TOTAL USD", sum(i['amount'] for i in invoices))
        
        if calculated_bf == 0:
            bf_label = f"BALANCE B/F {bf_date_str}"
        else:
            me_you = 'ME' if calculated_bf >= 0 else 'YOU'
            bf_label = f"BALANCE DUE TO {me_you} B/F ({bf_date_str})"
        add_summary_row(bf_label, calculated_bf)
        
        for tr in transfers:
            add_summary_row(f"TRANSFER RECEIVED {tr['date']}", tr['amount'])
            
        me_you_final = 'ME' if final_balance >= 0 else 'YOU'
        final_label = f"BALANCE DUE TO {me_you_final}"
        add_summary_row(final_label, final_balance)
        
        final_row_count = len(table.rows)
        row_height_cm = 1.15
        if final_row_count <= 15: row_height_cm = 1.15
        elif final_row_count == 16: row_height_cm = 1.07
        elif final_row_count == 17: row_height_cm = 1.01
        elif final_row_count == 18: row_height_cm = 0.95
        elif final_row_count == 19: row_height_cm = 0.89
        elif final_row_count == 20: row_height_cm = 0.83
        elif 21 <= final_row_count <= 35: row_height_cm = 1.15
        elif final_row_count == 36: row_height_cm = 1.12
        elif final_row_count == 37: row_height_cm = 1.11
        elif final_row_count == 38: row_height_cm = 1.07
        elif final_row_count == 39: row_height_cm = 1.04
        elif final_row_count == 40: row_height_cm = 1.02
        elif final_row_count in [41, 42]: row_height_cm = 0.98
        elif final_row_count >= 43: row_height_cm = 1.13
        
        for row in table.rows:
            row.height = Cm(row_height_cm)
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                
        docx_path = os.path.join(Config.STATEMENTS_DIR, f"{output_name}.docx")
        pdf_path = os.path.join(Config.STATEMENTS_DIR, f"{output_name}.pdf")
        
        doc.save(str(docx_path))
        
        pdf_msg = ""
        try:
            convert(str(docx_path), pdf_path)
            pdf_msg = f"{output_name} (Word + PDF) generated successfully."
        except Exception as pdf_err:
            pdf_msg = f"Warning: Word saved, but PDF conversion failed: {pdf_err}"
            
        if round(excel_expected_balance, 2) != round(final_balance, 2):
            pdf_msg += " | Warning: Balance doesn't match Excel expected balance."
            
        return True, pdf_msg
    except Exception as e:
        return False, f"Error generating statement: {e}"

def get_statement_rows():
    """Fetch current rows from the statement to display in the HTML interface"""
    wb = load_workbook(Config.EXCEL_PATH, data_only=True)
    ws = wb.active
    opening_rows = [r[0].row for r in ws.iter_rows(min_col=3, max_col=3) if str(r[0].value).strip() == Config.OPENING_TEXT]
    if not opening_rows:
        wb.close()
        return {"rows": [], "current_bf": 0}
        
    last_opening_row = opening_rows[-1]
    
    all_data = pd.read_excel(Config.EXCEL_PATH, header=None, engine='openpyxl')
    all_data[8] = pd.to_numeric(all_data[8], errors='coerce').fillna(0)
    
    sum_block_dif = 0
    for r in range(last_opening_row + 1, ws.max_row - 1):
        sum_block_dif += float(ws.cell(r, 9).value or 0)
        
    current_bf_balance = float(ws.cell(last_opening_row, 6).value or 0) + float(ws.cell(ws.max_row, 9).value) - sum_block_dif
    candidates = []
    running_balance = current_bf_balance
    
    for r in range(last_opening_row + 1, ws.max_row - 1):
        desc = ws.cell(r, 3).value or ""
        if not desc:
            continue
            
        edited_amt = ws.cell(r, 8).value
        actual_inv = ws.cell(r, 4).value
        payment_amt = ws.cell(r, 5).value
        
        row_amount = 0
        if "Attached Invoice" in desc:
            row_amount = float(edited_amt) if (edited_amt and edited_amt != 0) else float(actual_inv or 0)
            running_balance += row_amount
        elif "ADVANCE RECEIVED" in desc:
            row_amount = float(payment_amt or 0)
            running_balance -= row_amount
            
        date_val = ws.cell(r, 1).value
        candidates.append({
            "row_index": r,
            "date": date_val.strftime("%m/%d/%Y") if hasattr(date_val, 'strftime') else str(date_val),
            "inv_no": ws.cell(r, 2).value or "-",
            "desc": desc,
            "amount": row_amount,
            "running_balance": running_balance
        })
    wb.close()
    return {"rows": candidates, "current_bf": current_bf_balance}

def update_opening_balance(insert_at):
    """Insert a new opening balance row in Excel"""
    DESC_COL = 3
    try:
        wb = load_workbook(Config.EXCEL_PATH)
        ws = wb.active
        opening_rows = [r[0].row for r in ws.iter_rows(min_col=DESC_COL, max_col=DESC_COL) if str(r[0].value).strip() == Config.OPENING_TEXT]
        if not opening_rows:
            return False, "No opening row found."
            
        last_opening = opening_rows[-1]
        ws.insert_rows(insert_at, 3)
        ws.row_dimensions[insert_at].height = 18.75
        
        templates = [last_opening - 2, last_opening - 1, last_opening]
        for i in range(3):
            for col in range(1, ws.max_column + 1):
                src = ws.cell(templates[i], col)
                tgt = ws.cell(insert_at + i, col)
                if src.has_style:
                    tgt.font = copy(src.font)
                    tgt.border = copy(src.border)
                    tgt.fill = copy(src.fill)
                    tgt.alignment = copy(src.alignment)
                    tgt.number_format = src.number_format
                    tgt.protection = copy(src.protection)
                    
        total_row = insert_at + 1
        ws.cell(total_row, 4).value = f"=SUM(D{last_opening}:D{total_row-1})"
        ws.cell(total_row, 5).value = f"=SUM(E{last_opening}:E{total_row-1})"
        ws.cell(total_row, 6).value = f"=F{last_opening}+D{total_row}-E{total_row}"
        
        ob_row = insert_at + 2
        ws.cell(ob_row, DESC_COL).value = Config.OPENING_TEXT
        ws.cell(ob_row, 6).value = f"=F{total_row}"
        
        final_row = ws.max_row
        ws.cell(final_row, 4).value = f"=SUM(D{ob_row}:D{final_row-1})"
        ws.cell(final_row, 5).value = f"=SUM(E{ob_row}:E{final_row-1})"
        ws.cell(final_row, 6).value = f"=F{ob_row}+D{final_row}-E{final_row}"
        ws.cell(row=final_row, column=9).value = f"=SUM(I8:I{final_row-1})"
        ws.cell(row=final_row, column=11).value = f"=F{final_row}+I{final_row}"
        
        wb.save(Config.EXCEL_PATH)
        wb.close()
        refresh_excel_formulas()
        return True, "Opening balance updated successfully in Excel."
    except Exception as e:
        return False, f"Error inserting row: {e}"