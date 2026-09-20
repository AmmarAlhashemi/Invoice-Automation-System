import os
import re
import glob
import pandas as pd
from datetime import datetime
from copy import copy
import openpyxl
import xlwings as xw
from docxtpl import DocxTemplate, RichText
from num2words import num2words
from docx2pdf import convert
from config import Config

def refresh_excel_formulas():
    """Force Excel to recalculate all formulas in the background"""
    app = xw.App(visible=False)
    app.display_alerts = False
    try:
        wb = app.books.open(str(Config.EXCEL_PATH))
        wb.app.calculate()
        wb.save()
        wb.close()
    finally:
        app.quit()

def get_last_invoice_info(inv_type="normal"):
    """Fetch the latest invoice file based on modification date"""
    pattern = os.path.join(Config.INV_DIR, "INV-COC*.docx") if inv_type == "coc" else os.path.join(Config.INV_DIR, "INV-[0-9]*.docx")
    files = glob.glob(pattern)
    if not files:
        return 0, "No previous invoices"
    
    latest_file = max(files, key=os.path.getmtime)
    file_name = os.path.basename(latest_file)
    nums = re.findall(r'\d+', file_name)
    last_no = int(nums[0]) if nums else 0
    return last_no, file_name

def get_next_number_by_type(inv_type="normal"):
    """Fetch the next number for each invoice type"""
    pattern = os.path.join(Config.INV_DIR, "INV-COC*.docx") if inv_type == "coc" else os.path.join(Config.INV_DIR, "INV-[0-9]*.docx")
    files = glob.glob(pattern)
    if not files:
        return 1
    
    numbers = []
    for f in files:
        file_name = os.path.basename(f)
        nums = re.findall(r'\d+', file_name)
        if nums:
            numbers.append(int(nums[0]))
    return max(numbers) + 1 if numbers else 1

def update_excel_statement_logic(date_val, inv_no, amount, full_name):
    """Add the invoice to the Excel statement and update formulas"""
    if not os.path.exists(Config.EXCEL_PATH):
        return False, "Excel file not found!"
    
    try:
        if "COC" in inv_no:
            inv_no = inv_no.split('/')[0]
            
        wb = openpyxl.load_workbook(Config.EXCEL_PATH)
        ws = wb.active
        max_row = ws.max_row
        target_row = max_row - 1
        source_row = target_row - 2
        
        ws.insert_rows(target_row)
        ws.row_dimensions[ws.max_row - 1].height = 18.75
        
        for col in range(1, 7):
            new_cell = ws.cell(row=target_row, column=col)
            old_cell = ws.cell(row=source_row, column=col)
            if old_cell.has_style:
                new_cell.font = copy(old_cell.font)
                new_cell.border = copy(old_cell.border)
                new_cell.fill = copy(old_cell.fill)
                new_cell.number_format = copy(old_cell.number_format)
                new_cell.alignment = copy(old_cell.alignment)
                
        ws.cell(row=target_row, column=1).value = date_val
        ws.cell(row=target_row, column=2).value = inv_no
        ws.cell(row=target_row, column=3).value = "Attached Invoice"
        ws.cell(row=target_row, column=4).value = float(amount)
        
        max_row += 1  # Update formula reference
        for col in [4, 5, 9]:
            cell_val = ws.cell(row=max_row, column=col).value
            col_let = "D" if col == 4 else "E" if col == 5 else "I"
            if cell_val and ":" in cell_val:
                part1 = cell_val.split(':')[0]
                ws.cell(row=max_row, column=col).value = f"{part1}:{col_let}{max_row-1})"
                
        part1 = ws.cell(row=max_row, column=6).value.split('+')[0]
        ws.cell(row=max_row, column=6).value = f"{part1}+D{max_row}-E{max_row}"
        ws.cell(row=max_row, column=11).value = f"=F{max_row}+I{max_row}"
        
        wb.save(Config.EXCEL_PATH)
        wb.close()
        refresh_excel_formulas()
        return True, f"Excel statement updated successfully including {full_name}"
    except Exception as e:
        return False, f"Excel Error: {e}"

def process_and_generate_invoice(form_data):
    """Complete logic for processing and generating invoice files (Word + PDF)"""
    inv_type = form_data.get('inv_type_select')
    current_no = form_data.get('inv_no')
    year = datetime.now().year
    
    if inv_type == 'coc':
        clean_no = current_no.replace('/', '-')
        full_inv_name = f"INV-{clean_no}"
    else:
        full_inv_name = f"INV-{current_no}-{year}"
        
    docx_path = os.path.join(Config.INV_DIR, f"{full_inv_name}.docx")
    pdf_path = os.path.join(Config.INV_DIR, f"{full_inv_name}.pdf")
    
    if os.path.exists(docx_path):
        return False, f"No. ({current_no}) was taken by another invoice, Please choose another number!"
        
    try:
        doc = DocxTemplate(Config.INV_TEMPLATE_PATH)
        subtotal = 0.0
        context = {}
        
        for i in range(1, 11):
            context[f'd{i}'] = RichText(form_data.get(f'd{i}', ''))
            p_raw = form_data.get(f'p{i}', '0')
            if p_raw in ["00", "00.00"]:
                context[f'p{i}'] = "00.00"
            else:
                p_val = float(p_raw) if p_raw else 0.0
                subtotal += p_val
                context[f'p{i}'] = f"{p_val:,.2f}" if p_val > 0 else ""
                
        o_price = float(form_data.get('others_price', '0') or 0)
        subtotal += o_price
        
        context.update({
            'invoice_num': current_no if "COC" in current_no else f"{current_no}             ",
            'date_today': form_data.get('date'),
            'client_name': RichText(form_data.get('client_name', '')),
            'total_amount': f"{subtotal:,.2f}",
            'total_text': num2words(subtotal, lang='en').capitalize() + " dollars only.",
            'awb': RichText(form_data.get('awb', '')),
            'bol': RichText(form_data.get('bol', '')),
            'net_weight': form_data.get('net_w'),
            'gross_weight': form_data.get('gross_w'),
            'goods_nature': RichText(form_data.get('nature', '')),
            'd_others': RichText(form_data.get('others_desc', '')),
            'p_others': f"{o_price:,.2f}" if o_price > 0 else ""
        })
        
        doc.render(context)
        doc.save(docx_path)
        
        pdf_msg = ""
        try:
            convert(docx_path, pdf_path)
            pdf_msg = f"{full_inv_name} (Word + PDF) saved successfully."
        except Exception as pdf_err:
            pdf_msg = f"Warning: Word saved, but PDF failed. PDF Error: {pdf_err}"
            
        if 'update_excel' in form_data:
            excel_status, msg = update_excel_statement_logic(form_data.get('date'), current_no, subtotal, full_inv_name)
            return excel_status, f"{pdf_msg} | {msg}"
            
        return True, pdf_msg
    except Exception as e:
        return False, f"Error generating invoice: {e}"